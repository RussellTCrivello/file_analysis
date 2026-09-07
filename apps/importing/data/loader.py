"""
Data Loader
Loads domain data from Excel, CSV, or JSON files
"""

import json
import csv
from pathlib import Path
from typing import Dict, List, Optional
from abc import ABC, abstractmethod

from utils.logger import get_logger
from utils.exceptions import DataLoadException

logger = get_logger(__name__)

# Try to import openpyxl for Excel support
try:
    from openpyxl import load_workbook
    EXCEL_SUPPORT = True
except ImportError:
    EXCEL_SUPPORT = False
    logger.warning("openpyxl not available - Excel support disabled")


class DataLoader(ABC):
    """Abstract base class for data loaders"""
    
    @abstractmethod
    def load(self, file_path: Path) -> Dict[str, List[str]]:
        """Load data from file"""
        pass
    
    @abstractmethod
    def can_handle(self, file_path: Path) -> bool:
        """Check if this loader can handle the file"""
        pass


class ExcelLoader(DataLoader):
    """Loads domain data from Excel files"""
    
    def can_handle(self, file_path: Path) -> bool:
        """Check if file is Excel format"""
        return file_path.suffix.lower() == '.xlsx' and EXCEL_SUPPORT
    
    def load(self, file_path: Path) -> Dict[str, List[str]]:
        """
        Load domain data from Excel file
        
        Structure:
        - Each sheet = One domain
        - Sheet name = Domain name
        - Words can be in any cells
        """
        if not EXCEL_SUPPORT:
            raise DataLoadException("Excel support not available. Install: pip install openpyxl")
        
        # Resolve path to absolute if it's relative
        resolved_path = file_path.resolve() if not file_path.is_absolute() else file_path
        
        if not resolved_path.exists():
            raise DataLoadException(
                f"Excel file not found: {file_path}\n"
                f"Resolved path: {resolved_path}\n"
                f"Current working directory: {Path.cwd()}"
            )
        
        try:
            wb = load_workbook(resolved_path, read_only=True, data_only=True)
            domains = {}
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                words = []
                
                # Read all cells in the sheet
                for row in sheet.iter_rows(values_only=True):
                    for cell_value in row:
                        if cell_value is not None:
                            word = str(cell_value).strip()
                            if word:
                                words.append(word)
                
                if words:
                    # Remove duplicates while preserving order
                    unique_words = list(dict.fromkeys(words))
                    domains[sheet_name] = unique_words
                    logger.debug(f"Loaded {len(unique_words)} words for '{sheet_name}'")
                else:
                    logger.warning(f"Sheet '{sheet_name}' contains no words")
            
            wb.close()
            
            if not domains:
                raise DataLoadException(f"No domain data found in {resolved_path}")
            
            logger.info(f"Loaded {len(domains)} domains from Excel")
            return domains
            
        except DataLoadException:
            raise
        except Exception as e:
            raise DataLoadException(f"Failed to load Excel file {resolved_path}: {e}") from e


class CSVLoader(DataLoader):
    """Loads domain data from CSV files"""
    
    def can_handle(self, file_path: Path) -> bool:
        """Check if file is CSV format"""
        return file_path.suffix.lower() == '.csv'
    
    def load(self, file_path: Path) -> Dict[str, List[str]]:
        """
        Load domain data from CSV file
        
        Format:
        domain_name, word1, word2, word3, ...
        """
        if not file_path.exists():
            raise DataLoadException(f"CSV file not found: {file_path}")
        
        try:
            domains = {}
            current_domain = None
            current_words = []
            
            with open(file_path, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                
                for row in reader:
                    if not row or not any(row):
                        continue
                    
                    row = [cell.strip() for cell in row if cell.strip()]
                    if not row:
                        continue
                    
                    domain_name = row[0]
                    words = row[1:] if len(row) > 1 else []
                    
                    if domain_name == current_domain:
                        current_words.extend(words)
                    else:
                        if current_domain and current_words:
                            domains[current_domain] = current_words
                        current_domain = domain_name
                        current_words = words.copy()
                
                # Don't forget last domain
                if current_domain and current_words:
                    # Remove duplicates while preserving order
                    unique_words = list(dict.fromkeys(current_words))
                    domains[current_domain] = unique_words
            
            # Remove duplicates from all domains
            for domain_name in domains:
                # Remove duplicates while preserving order
                domains[domain_name] = list(dict.fromkeys(domains[domain_name]))
            
            if not domains:
                raise DataLoadException(f"No domain data found in {file_path}")
            
            logger.info(f"Loaded {len(domains)} domains from CSV")
            return domains
            
        except Exception as e:
            raise DataLoadException(f"Failed to load CSV file: {e}") from e


class JSONLoader(DataLoader):
    """Loads domain data from JSON files"""
    
    def can_handle(self, file_path: Path) -> bool:
        """Check if file is JSON format"""
        return file_path.suffix.lower() == '.json'
    
    def load(self, file_path: Path) -> Dict[str, List[str]]:
        """
        Load domain data from JSON file
        
        Format:
        {
            "domain1": ["word1", "word2", ...],
            "domain2": ["word3", "word4", ...]
        }
        """
        if not file_path.exists():
            raise DataLoadException(f"JSON file not found: {file_path}")
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                domains = json.load(f)
            
            if not isinstance(domains, dict):
                raise DataLoadException("JSON must be a dictionary")
            
            if not domains:
                raise DataLoadException(f"No domain data found in {file_path}")
            
            logger.info(f"Loaded {len(domains)} domains from JSON")
            return domains
            
        except json.JSONDecodeError as e:
            raise DataLoadException(f"Invalid JSON format: {e}") from e
        except Exception as e:
            raise DataLoadException(f"Failed to load JSON file: {e}") from e


class DataLoaderFactory:
    """Factory for creating appropriate data loader"""
    
    def __init__(self):
        self.loaders = [
            ExcelLoader(),
            CSVLoader(),
            JSONLoader(),
        ]
    
    def get_loader(self, file_path: Path) -> DataLoader:
        """Get appropriate loader for file"""
        for loader in self.loaders:
            if loader.can_handle(file_path):
                return loader
        
        raise DataLoadException(
            f"No loader available for {file_path.suffix}. "
            f"Supported: .xlsx, .csv, .json"
        )
    
    def load(self, file_path: str) -> Dict[str, List[str]]:
        """Load data from file using appropriate loader"""
        path = Path(file_path)
        loader = self.get_loader(path)
        logger.info(f"Using {loader.__class__.__name__} for {path}")
        return loader.load(path)
